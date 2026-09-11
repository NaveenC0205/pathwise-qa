"""Read CSV/XLSX test suites without evaluating spreadsheet formulas."""
import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from .models import ExistingCase

ALIASES = {
    'name': ['name', 'title', 'test case name', 'test case title', 'testcase', 'test case', 'test scenario', 'scenario', 'summary'],
    'description': ['description', 'test description', 'test case description', 'objective'],
    'module': ['module', 'module name', 'feature', 'component'],
    'steps': ['steps', 'test steps', 'test step', 'test step description', 'action', 'actions', 'procedure'],
    'expected_result': ['expected result', 'expected results', 'expected', 'expected outcome', 'step expected result'],
    'testing_type': ['testing type', 'test type', 'type', 'kind'],
    'suite': ['suite', 'test suite', 'test category'],
}


def normalize(value):
    return re.sub(r'\s+', ' ', re.sub(r'[_\-]+', ' ', str(value or '').strip().lower()))


def import_cases(data: bytes, filename: str):
    if len(data) > 2_000_000:
        raise ValueError('Use a CSV or XLSX file smaller than 2 MB.')
    suffix = Path(filename).suffix.lower()
    sheets = []
    warnings = []
    if suffix == '.csv':
        try: content = data.decode('utf-8-sig')
        except UnicodeDecodeError as exc: raise ValueError('Save the CSV as UTF-8 before importing.') from exc
        try: dialect = csv.Sniffer().sniff(content[:8192], delimiters=',;\t')
        except csv.Error: dialect = csv.excel
        sheets = [('CSV', csv.reader(io.StringIO(content), dialect))]
    elif suffix == '.xlsx':
        try:
            archive = zipfile.ZipFile(io.BytesIO(data))
            if len(archive.infolist()) > 500 or sum(f.file_size for f in archive.infolist()) > 15_000_000:
                raise ValueError('Workbook is too large after decompression.')
            archive.close()
        except zipfile.BadZipFile as exc: raise ValueError('This is not a valid XLSX workbook.') from exc
        from openpyxl import load_workbook
        try: workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=False, keep_links=False)
        except Exception as exc: raise ValueError('The workbook could not be read. Use an unencrypted .xlsx file.') from exc
        if len(workbook.worksheets) > 20:
            workbook.close(); raise ValueError('Use at most 20 worksheets per import.')
        sheets = [(sheet.title, sheet.iter_rows(values_only=True)) for sheet in workbook.worksheets]
    else:
        raise ValueError('Supported formats are .csv and .xlsx. Save older .xls files as .xlsx.')
    cases = []
    fingerprint = hashlib.sha256(data).hexdigest()[:10]
    total_chars = 0
    try:
        for sheet_name, rows in sheets:
            mapping = None
            for row_no, row in enumerate(rows, 1):
                if row_no > 2000: raise ValueError('Use at most 2,000 rows per sheet.')
                if len(row) > 100: raise ValueError('Use at most 100 columns per sheet.')
                values = [str(v) if v is not None else '' for v in row]
                if not any(v.strip() for v in values): continue
                if mapping is None:
                    header = [normalize(v) for v in values]
                    mapping = {key: next((i for i, v in enumerate(header) if v in aliases), None) for key, aliases in ALIASES.items()}
                    if mapping['name'] is None:
                        raise ValueError(f'{sheet_name}: add a Test Case Name, Name, Title, or Summary header in the first nonempty row.')
                    for col in ('steps', 'expected_result'):
                        if mapping[col] is None: warnings.append(f'{sheet_name}: no {col.replace("_", " ")} column; the agents will flag incomplete cases.')
                    continue
                fields = {key: values[index].strip() if index is not None and index < len(values) else '' for key, index in mapping.items()}
                # Formulas are not executed or used as test instructions.
                if any(value.startswith('=') for value in fields.values()):
                    warnings.append(f'{sheet_name} row {row_no}: formula cells were excluded.')
                    fields = {k: '' if v.startswith('=') else v for k, v in fields.items()}
                if not fields['name']:
                    warnings.append(f'{sheet_name} row {row_no}: missing case name; row was not imported.'); continue
                total_chars += sum(map(len, fields.values()))
                if total_chars > 140000: raise ValueError('Imported test content exceeds 140,000 characters; split the file by module.')
                row_id = hashlib.sha256(f'{fingerprint}:{sheet_name}:{row_no}'.encode()).hexdigest()[:14]
                cases.append(ExistingCase(id=f'EXT-{row_id}', source_file=Path(filename).name, source_sheet=sheet_name, source_row=row_no, **fields))
                if len(cases) > 200: raise ValueError('Import at most 200 named cases per file.')
    finally:
        if suffix == '.xlsx': workbook.close()
    if not cases: raise ValueError('No named test cases found in the file.')
    return {'cases': [c.model_dump() for c in cases], 'warnings': warnings, 'count': len(cases), 'filename': Path(filename).name}
